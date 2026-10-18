"""
Q-04 import_store_data.py 검증 pytest.

정상 파일 1개 + 오류 파일 3종 픽스처를 tmp_path에 생성해 테스트합니다.
PostgreSQL 불필요 — openpyxl + 검증 로직만 검사합니다.
"""
import sys
from pathlib import Path

import openpyxl
import pytest

# scripts 폴더를 import 경로에 추가
sys.path.insert(0, str(Path(__file__).parent.parent))

from import_store_data import (
    StoreDataImporter,
    StoreDataValidationError,
    StoreDataValidator,
)


# ──────────────────────────────────────────────────────────────
# 픽스처 헬퍼
# ──────────────────────────────────────────────────────────────

def _make_workbook(
    menus: list[tuple],
    ingredients: list[tuple],
    recipes: list[tuple],
) -> openpyxl.Workbook:
    """헤더 + 데이터 행으로 워크북 생성."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_m = wb.create_sheet("메뉴")
    ws_m.append(["메뉴명", "가격(원)", "카테고리", "판매중(Y/N)"])
    for row in menus:
        ws_m.append(list(row))

    ws_i = wb.create_sheet("식재료")
    ws_i.append(["식재료명", "단위", "단가(원)", "현재고", "안전재고", "리드타임(일)"])
    for row in ingredients:
        ws_i.append(list(row))

    ws_r = wb.create_sheet("레시피")
    ws_r.append(["메뉴명", "식재료명", "수량(qty_per_unit)", "단위"])
    for row in recipes:
        ws_r.append(list(row))

    return wb


def _save(wb, path: Path) -> Path:
    wb.save(str(path))
    return path


# ──────────────────────────────────────────────────────────────
# 정상 파일 픽스처
# ──────────────────────────────────────────────────────────────

@pytest.fixture
def normal_xlsx(tmp_path: Path) -> Path:
    """정상 파일: 메뉴1 + 식재료1 + 레시피1."""
    wb = _make_workbook(
        menus=[("아메리카노", 4500, "drink", "Y")],
        ingredients=[("원두", "g", 30, 5000, 1000, 3)],
        recipes=[("아메리카노", "원두", 15, "g")],
    )
    return _save(wb, tmp_path / "normal.xlsx")


# ──────────────────────────────────────────────────────────────
# 오류 파일 3종 픽스처
# ──────────────────────────────────────────────────────────────

@pytest.fixture
def err_missing_required_xlsx(tmp_path: Path) -> Path:
    """오류 1: 메뉴 시트에서 메뉴명 필수 칸 비어있음 (2행)."""
    wb = _make_workbook(
        menus=[("", 4500, "drink", "Y")],           # 메뉴명 공백
        ingredients=[("원두", "g", 30, 5000, 1000, 3)],
        recipes=[("아메리카노", "원두", 15, "g")],
    )
    return _save(wb, tmp_path / "err_missing_required.xlsx")


@pytest.fixture
def err_invalid_ref_xlsx(tmp_path: Path) -> Path:
    """오류 2: 레시피에서 식재료 시트에 없는 식재료명 참조."""
    wb = _make_workbook(
        menus=[("아메리카노", 4500, "drink", "Y")],
        ingredients=[("원두", "g", 30, 5000, 1000, 3)],
        recipes=[("아메리카노", "없는재료", 15, "g")],  # 식재료 시트에 없음
    )
    return _save(wb, tmp_path / "err_invalid_ref.xlsx")


@pytest.fixture
def err_invalid_number_xlsx(tmp_path: Path) -> Path:
    """오류 3: 가격(원) 칸에 문자열 입력."""
    wb = _make_workbook(
        menus=[("아메리카노", "사천오백원", "drink", "Y")],  # 숫자여야 함
        ingredients=[("원두", "g", 30, 5000, 1000, 3)],
        recipes=[("아메리카노", "원두", 15, "g")],
    )
    return _save(wb, tmp_path / "err_invalid_number.xlsx")


# ──────────────────────────────────────────────────────────────
# 테스트
# ──────────────────────────────────────────────────────────────

class TestStoreDataValidator:

    def test_normal_file_has_no_issues(self, normal_xlsx: Path):
        """정상 파일: 검증 오류 0건."""
        wb = openpyxl.load_workbook(str(normal_xlsx), read_only=True, data_only=True)
        issues = StoreDataValidator().validate(wb)
        assert issues == [], f"정상 파일에서 오류 발생: {issues}"

    def test_missing_required_field_detected(self, err_missing_required_xlsx: Path):
        """오류 1: 필수 칸 누락 → 해당 행·컬럼 ValidationIssue 포함."""
        wb = openpyxl.load_workbook(str(err_missing_required_xlsx), read_only=True, data_only=True)
        issues = StoreDataValidator().validate(wb)
        assert any("메뉴명" in i.message and i.row == 2 for i in issues), (
            f"필수 칸 누락 미탐지: {issues}"
        )

    def test_invalid_ingredient_reference_detected(self, err_invalid_ref_xlsx: Path):
        """오류 2: 레시피에서 없는 식재료명 참조 → 참조 오류 포함."""
        wb = openpyxl.load_workbook(str(err_invalid_ref_xlsx), read_only=True, data_only=True)
        issues = StoreDataValidator().validate(wb)
        assert any("없는재료" in i.message for i in issues), (
            f"참조 오류 미탐지: {issues}"
        )

    def test_invalid_number_field_detected(self, err_invalid_number_xlsx: Path):
        """오류 3: 숫자 칸에 문자열 → 숫자 오류 포함."""
        wb = openpyxl.load_workbook(str(err_invalid_number_xlsx), read_only=True, data_only=True)
        issues = StoreDataValidator().validate(wb)
        assert any("가격(원)" in i.message for i in issues), (
            f"숫자 오류 미탐지: {issues}"
        )


class TestStoreDataImporter:

    def test_normal_file_loads_and_returns_counts(self, normal_xlsx: Path):
        """정상 파일: 적재 성공, 반환값 확인."""
        result = StoreDataImporter().load_file(normal_xlsx)
        assert result["menus"] == 1
        assert result["ingredients"] == 1
        assert result["recipes"] == 1

    def test_idempotent_double_load(self, normal_xlsx: Path):
        """같은 파일 두 번 올려도 예외 없음 (upsert)."""
        importer = StoreDataImporter()
        r1 = importer.load_file(normal_xlsx)
        r2 = importer.load_file(normal_xlsx)
        assert r1 == r2, "두 번째 적재 결과가 다릅니다 (중복 생성 의심)"

    def test_validation_error_raised_on_invalid_file(
        self, err_missing_required_xlsx: Path
    ):
        """오류 파일: StoreDataValidationError 발생."""
        with pytest.raises(StoreDataValidationError) as exc_info:
            StoreDataImporter().load_file(err_missing_required_xlsx)
        assert exc_info.value.issues, "issues 리스트가 비어 있습니다"
