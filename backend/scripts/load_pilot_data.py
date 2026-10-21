"""파일럿 초기 데이터 적재 스크립트 (B-19).

`private/pilot-store-data.xlsx`의 메뉴·식재료를 공개 API로 넣는다. DB에 직접 접속하지 않는다.
- 사장으로 로그인(POST /auth/login)한 뒤 /menus, /ingredients, /ingredients/{id}/movements를 쓴다.
- 이름이 같은 항목은 새로 만들지 않고 값이 다를 때만 고친다. 그래서 여러 번 실행해도 된다.
  (--create-only를 주면 이미 있는 항목은 건드리지 않는다. 사장이 실제 가격으로 고친 뒤 다시 실행할 때 쓴다.)
- 메뉴 종류가 main/side가 아닌 항목(주류)은 일시품절(is_active=false)로 넣는다. 판매가가 가상 값이기 때문이다.
- 시작 재고는 새로 만든 식재료(또는 재고가 0인 기존 식재료)에만 입고(restock) 기록으로 넣는다.
- xlsx 파일과 그 값은 저장소에 두지 않는다. 비밀번호는 환경 변수(POS_OWNER_PASSWORD)나 입력창으로만 받는다.

실행 (backend 폴더에서, openpyxl 필요):
    python scripts/load_pilot_data.py --base-url https://<주소> --xlsx <파일> --dry-run
    python scripts/load_pilot_data.py --base-url https://<주소> --xlsx <파일>
"""
import argparse
import getpass
import os
import sys
from dataclasses import dataclass, field
from pathlib import Path

SELLABLE_CATEGORIES = ("main", "side")
API_PREFIX = "/api/v1"
MENU_SHEET = "메뉴"
INGREDIENT_SHEET = "식재료"
INITIAL_STOCK_TYPE = "restock"


class LoadError(Exception):
    """입력 파일 또는 API 응답 문제. 메시지만 출력하고 종료한다(토큰·비밀번호는 담지 않는다)."""


@dataclass(frozen=True)
class MenuRow:
    name: str
    price: float
    category: str | None
    is_active: bool


@dataclass(frozen=True)
class IngredientRow:
    name: str
    unit: str
    unit_price: float
    current_stock: float
    safety_stock: float
    lead_time_days: int


@dataclass
class Report:
    created: dict = field(default_factory=lambda: {"menus": 0, "ingredients": 0, "stock_movements": 0})
    updated: dict = field(default_factory=lambda: {"menus": 0, "ingredients": 0})
    unchanged: dict = field(default_factory=lambda: {"menus": 0, "ingredients": 0})


# ── 파일 읽기 (검증만 하고 API는 부르지 않는다) ─────────────────────────────

def _num(value, label: str, row_no: int, *, integer: bool = False) -> float:
    if value is None or value == "":
        return 0
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise LoadError(f"{label} {row_no}행: 숫자가 아닙니다") from None
    if number < 0:
        raise LoadError(f"{label} {row_no}행: 음수입니다")
    if integer and number != int(number):
        raise LoadError(f"{label} {row_no}행: 정수여야 합니다")
    # 서버는 금액·수량을 소수 둘째 자리까지 저장한다. 같은 자리로 맞춰야 다시 실행했을 때 '다름'으로 보이지 않는다.
    return int(number) if integer else round(number, 2)


def parse_menu_rows(rows) -> list[MenuRow]:
    """rows: 헤더를 뺀 (이름, 가격, 종류, 판매여부 Y/N) 튜플들."""
    result, seen = [], set()
    for row_no, row in enumerate(rows, start=2):
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        name = str(row[0]).strip()
        if name in seen:
            raise LoadError(f"{MENU_SHEET} {row_no}행: 이름이 겹칩니다 ({name})")
        seen.add(name)
        if row[1] is None:
            raise LoadError(f"{MENU_SHEET} {row_no}행: 판매 가격이 없습니다 ({name})")
        price = _num(row[1], MENU_SHEET, row_no)
        category = str(row[2]).strip() if len(row) > 2 and row[2] not in (None, "") else None
        flag = str(row[3]).strip().upper() if len(row) > 3 and row[3] is not None else ""
        if flag not in ("Y", "N"):
            raise LoadError(f"{MENU_SHEET} {row_no}행: 판매 여부는 Y 또는 N이어야 합니다 ({name})")
        # 주류 등 main/side 밖의 종류는 가격이 확정될 때까지 손님 메뉴판에 올리지 않는다.
        is_active = flag == "Y" and category in SELLABLE_CATEGORIES
        result.append(MenuRow(name=name, price=price, category=category, is_active=is_active))
    return result


def parse_ingredient_rows(rows) -> list[IngredientRow]:
    """rows: 헤더를 뺀 (이름, 단위, 단가, 현재고, 안전재고, 리드타임) 튜플들."""
    result, seen = [], set()
    for row_no, row in enumerate(rows, start=2):
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        name = str(row[0]).strip()
        if name in seen:
            raise LoadError(f"{INGREDIENT_SHEET} {row_no}행: 이름이 겹칩니다 ({name})")
        seen.add(name)
        unit = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if not unit:
            raise LoadError(f"{INGREDIENT_SHEET} {row_no}행: 단위가 없습니다 ({name})")
        cells = list(row) + [None] * (6 - len(row))
        result.append(
            IngredientRow(
                name=name,
                unit=unit,
                unit_price=_num(cells[2], INGREDIENT_SHEET, row_no),
                current_stock=_num(cells[3], INGREDIENT_SHEET, row_no),
                safety_stock=_num(cells[4], INGREDIENT_SHEET, row_no),
                lead_time_days=_num(cells[5], INGREDIENT_SHEET, row_no, integer=True),
            )
        )
    return result


def read_workbook(path: Path) -> tuple[list[MenuRow], list[IngredientRow]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise LoadError("openpyxl이 필요합니다: pip install openpyxl") from None
    if not Path(path).is_file():
        raise LoadError(f"파일이 없습니다: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in (MENU_SHEET, INGREDIENT_SHEET):
        if sheet not in wb.sheetnames:
            raise LoadError(f"'{sheet}' 시트가 없습니다")
    menus = parse_menu_rows(list(wb[MENU_SHEET].iter_rows(min_row=2, values_only=True)))
    ingredients = parse_ingredient_rows(list(wb[INGREDIENT_SHEET].iter_rows(min_row=2, values_only=True)))
    return menus, ingredients


# ── API 호출 (client는 httpx.Client 호환: 상대 경로 사용) ───────────────────

def _check(response, action: str):
    if response.status_code >= 400:
        try:
            detail = response.json()
        except ValueError:
            detail = response.text[:200]
        raise LoadError(f"{action} 실패 (HTTP {response.status_code}): {detail}")
    return response


def login(client, username: str, password: str) -> str:
    r = _check(client.post("/auth/login", json={"username": username, "password": password}), "로그인")
    return r.json()["access_token"]


def _list_items(client, path: str, headers: dict) -> list[dict]:
    items, cursor = [], None
    while True:
        params = {"limit": 200, **({"cursor": cursor} if cursor else {})}
        body = _check(client.get(path, params=params, headers=headers), f"GET {path}").json()
        items.extend(body["items"])
        cursor = body.get("next_cursor")
        if not body.get("has_more") or not cursor:
            return items


def load_menus(client, headers: dict, menus: list[MenuRow], report: Report, *, dry_run: bool, create_only: bool) -> None:
    existing = {m["name"]: m for m in _list_items(client, "/menus", headers)}
    for row in menus:
        payload = {"name": row.name, "price": row.price, "category": row.category, "is_active": row.is_active}
        current = existing.get(row.name)
        if current is None:
            if not dry_run:
                _check(client.post("/menus", json=payload, headers=headers), f"메뉴 생성({row.name})")
            report.created["menus"] += 1
        elif create_only or (
            float(current["price"]) == row.price
            and current.get("category") == row.category
            and bool(current["is_active"]) == row.is_active
        ):
            report.unchanged["menus"] += 1
        else:
            if not dry_run:
                _check(client.patch(f"/menus/{current['id']}", json=payload, headers=headers), f"메뉴 수정({row.name})")
            report.updated["menus"] += 1


def _add_initial_stock(client, headers: dict, ingredient_id: str, row: IngredientRow, report: Report, dry_run: bool) -> None:
    if row.current_stock <= 0:
        return
    if not dry_run:
        _check(
            client.post(
                f"/ingredients/{ingredient_id}/movements",
                json={"qty_delta": row.current_stock, "type": INITIAL_STOCK_TYPE},
                headers=headers,
            ),
            f"시작 재고 입고({row.name})",
        )
    report.created["stock_movements"] += 1


def load_ingredients(
    client, headers: dict, ingredients: list[IngredientRow], report: Report, *, dry_run: bool, create_only: bool
) -> None:
    existing = {i["name"]: i for i in _list_items(client, "/ingredients", headers)}
    for row in ingredients:
        payload = {
            "name": row.name,
            "unit": row.unit,
            "unit_price": row.unit_price,
            "safety_stock": row.safety_stock,
            "lead_time_days": row.lead_time_days,
        }
        current = existing.get(row.name)
        if current is None:
            ingredient_id = None
            if not dry_run:
                ingredient_id = _check(
                    client.post("/ingredients", json=payload, headers=headers), f"식재료 등록({row.name})"
                ).json()["id"]
            report.created["ingredients"] += 1
            _add_initial_stock(client, headers, ingredient_id, row, report, dry_run)
            continue
        if create_only:
            report.unchanged["ingredients"] += 1
            continue
        same = (
            current["unit"] == row.unit
            and float(current["unit_price"]) == row.unit_price
            and float(current["safety_stock"]) == row.safety_stock
            and int(current["lead_time_days"]) == row.lead_time_days
        )
        if same:
            report.unchanged["ingredients"] += 1
        else:
            if not dry_run:
                _check(
                    client.patch(f"/ingredients/{current['id']}", json=payload, headers=headers),
                    f"식재료 수정({row.name})",
                )
            report.updated["ingredients"] += 1
        # 이전 실행이 재고 입고 전에 끊긴 경우만 이어서 넣는다(재고가 있으면 건드리지 않는다).
        if float(current["current_stock"]) == 0:
            _add_initial_stock(client, headers, current["id"], row, report, dry_run)


def run(client, token: str, menus, ingredients, *, dry_run: bool = False, create_only: bool = False) -> Report:
    headers = {"Authorization": f"Bearer {token}"}
    report = Report()
    load_menus(client, headers, menus, report, dry_run=dry_run, create_only=create_only)
    load_ingredients(client, headers, ingredients, report, dry_run=dry_run, create_only=create_only)
    return report


def format_report(report: Report, menus, ingredients, *, dry_run: bool) -> str:
    inactive = sum(1 for m in menus if not m.is_active)
    lines = [
        ("[dry-run: 아무것도 쓰지 않았습니다] " if dry_run else "") + f"파일: 메뉴 {len(menus)}개(일시품절 {inactive}), 식재료 {len(ingredients)}개",
        f"새로 만듦: {report.created}",
        f"고침: {report.updated}",
        f"그대로: {report.unchanged}",
    ]
    return "\n".join(lines)


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="파일럿 메뉴·식재료 적재 (공개 API 사용)")
    p.add_argument("--base-url", required=True, help="서버 주소 (예: http://localhost:8000). /api/v1은 자동으로 붙인다")
    p.add_argument("--xlsx", required=True, type=Path, help="private/pilot-store-data.xlsx 경로")
    p.add_argument("--username", default="owner", help="사장 로그인 아이디(기본 owner)")
    p.add_argument("--dry-run", action="store_true", help="읽기만 하고 무엇을 바꿀지 요약만 출력")
    p.add_argument("--create-only", action="store_true", help="이미 있는 항목은 고치지 않고 없는 것만 만든다")
    return p


def _api_url(base_url: str) -> str:
    base = base_url.rstrip("/")
    return base if base.endswith(API_PREFIX) else base + API_PREFIX


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        menus, ingredients = read_workbook(args.xlsx)  # 파일 오류는 서버를 건드리기 전에 걸러낸다
        password = os.environ.get("POS_OWNER_PASSWORD") or getpass.getpass("사장 비밀번호: ")
        import httpx

        with httpx.Client(base_url=_api_url(args.base_url), timeout=30) as client:
            token = login(client, args.username, password)
            report = run(client, token, menus, ingredients, dry_run=args.dry_run, create_only=args.create_only)
    except LoadError as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 1
    print(format_report(report, menus, ingredients, dry_run=args.dry_run))
    return 0


if __name__ == "__main__":
    sys.exit(main())
