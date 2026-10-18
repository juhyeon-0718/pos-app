"""
Q-04 매장 데이터 적재 스크립트.

엑셀 파일(store-data-template.xlsx)을 읽어 검증 후
B-06 공개 service로 DB에 적재합니다.

B-06 머지 전: _menu_upsert·_ingredient_upsert·_recipe_upsert는 stub.
B-06 머지 후: app.menu.service, app.inventory.service로 교체.

실매장 엑셀은 커밋하지 않습니다(D-28).
"""
from __future__ import annotations

import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl이 설치되지 않았습니다: pip install openpyxl")


# ──────────────────────────────────────────────────────────────
# 검증
# ──────────────────────────────────────────────────────────────

@dataclass
class ValidationIssue:
    sheet: str
    row: int        # 1-based (헤더=1, 데이터=2~)
    col: str        # 컬럼 레터 또는 ""
    message: str

    def __str__(self) -> str:
        loc = f"[{self.sheet}!{self.col}{self.row}]" if self.col else f"[{self.sheet} 행{self.row}]"
        return f"{loc} {self.message}"


class StoreDataValidationError(Exception):
    def __init__(self, issues: list[ValidationIssue]) -> None:
        self.issues = issues
        lines = "\n".join(f"  - {i}" for i in issues)
        super().__init__(f"{len(issues)}건의 검증 오류:\n{lines}")


class StoreDataValidator:
    """엑셀 워크북 검증. validate() → issues 리스트 반환 (빈 리스트 = 정상)."""

    # 시트별 필수 헤더 (헤더 파싱 후 ' *' 제거된 이름 기준)
    MENU_REQUIRED: frozenset[str] = frozenset({"메뉴명", "가격(원)", "판매중(Y/N)"})
    INGREDIENT_REQUIRED: frozenset[str] = frozenset({"식재료명", "단위", "현재고"})
    RECIPE_REQUIRED: frozenset[str] = frozenset({"메뉴명", "식재료명", "수량(qty_per_unit)", "단위"})

    # 시트별 숫자여야 하는 헤더
    NUMERIC: dict[str, frozenset[str]] = {
        "메뉴": frozenset({"가격(원)"}),
        "식재료": frozenset({"단가(원)", "현재고", "안전재고", "리드타임(일)"}),
        "레시피": frozenset({"수량(qty_per_unit)"}),
    }

    # ── 공개 API ──────────────────────────────────────────────

    def validate(self, wb) -> list[ValidationIssue]:
        issues: list[ValidationIssue] = []

        for name in ("메뉴", "식재료", "레시피"):
            if name not in wb.sheetnames:
                issues.append(ValidationIssue(name, 0, "", f"'{name}' 시트가 없습니다"))

        if issues:
            return issues

        menu_names = self._validate_data_sheet(
            wb["메뉴"], self.MENU_REQUIRED, self.NUMERIC["메뉴"], issues, name_col="메뉴명"
        )
        ingredient_names = self._validate_data_sheet(
            wb["식재료"], self.INGREDIENT_REQUIRED, self.NUMERIC["식재료"], issues, name_col="식재료명"
        )
        self._validate_recipe_sheet(wb["레시피"], menu_names, ingredient_names, issues)
        return issues

    # ── 내부 헬퍼 ─────────────────────────────────────────────

    def parse_header(self, ws) -> dict[str, str]:
        """첫 행 헤더 → {정규화된 이름: 컬럼 레터} 매핑."""
        result: dict[str, str] = {}
        for cell in ws[1]:
            if cell.value is not None:
                name = str(cell.value).replace(" *", "").strip()
                result[name] = cell.column_letter
        return result

    def _row_dict(self, row: tuple, header: dict[str, str]) -> dict[str, Any]:
        d: dict[str, Any] = {}
        for col_name, col_letter in header.items():
            idx = ord(col_letter) - ord("A")
            d[col_name] = row[idx] if idx < len(row) else None
        return d

    def _is_blank(self, val: Any) -> bool:
        return val is None or (isinstance(val, str) and not val.strip())

    def _is_empty_row(self, row: tuple) -> bool:
        return all(self._is_blank(v) for v in row)

    def _validate_data_sheet(
        self,
        ws,
        required: frozenset[str],
        numeric: frozenset[str],
        issues: list[ValidationIssue],
        name_col: str,
    ) -> set[str]:
        """시트 검증. 유효한 이름(name_col 값) 집합 반환."""
        header = self.parse_header(ws)
        sheet = ws.title

        for col in required:
            if col not in header:
                issues.append(ValidationIssue(sheet, 1, "", f"필수 헤더 '{col}'이 없습니다"))

        names: set[str] = set()

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if self._is_empty_row(row):
                continue

            rd = self._row_dict(row, header)

            for col in required:
                if col in header and self._is_blank(rd.get(col)):
                    issues.append(ValidationIssue(sheet, row_idx, header[col], f"필수 칸 '{col}'이 비어 있습니다"))

            for col in numeric:
                if col not in header:
                    continue
                val = rd.get(col)
                if not self._is_blank(val):
                    try:
                        float(val)  # type: ignore[arg-type]
                    except (TypeError, ValueError):
                        issues.append(ValidationIssue(
                            sheet, row_idx, header[col],
                            f"'{col}' 칸에 숫자가 아닌 값 '{val}'이 있습니다",
                        ))

            name_val = rd.get(name_col)
            if name_val and not self._is_blank(name_val):
                names.add(str(name_val).strip())

        return names

    def _validate_recipe_sheet(
        self,
        ws,
        menu_names: set[str],
        ingredient_names: set[str],
        issues: list[ValidationIssue],
    ) -> None:
        header = self.parse_header(ws)
        sheet = ws.title

        for col in self.RECIPE_REQUIRED:
            if col not in header:
                issues.append(ValidationIssue(sheet, 1, "", f"필수 헤더 '{col}'이 없습니다"))

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if self._is_empty_row(row):
                continue

            rd = self._row_dict(row, header)

            for col in self.RECIPE_REQUIRED:
                if col in header and self._is_blank(rd.get(col)):
                    issues.append(ValidationIssue(sheet, row_idx, header[col], f"필수 칸 '{col}'이 비어 있습니다"))

            # 숫자 검증
            for col in self.NUMERIC["레시피"]:
                if col not in header:
                    continue
                val = rd.get(col)
                if not self._is_blank(val):
                    try:
                        float(val)  # type: ignore[arg-type]
                    except (TypeError, ValueError):
                        issues.append(ValidationIssue(
                            sheet, row_idx, header[col],
                            f"'{col}' 칸에 숫자가 아닌 값 '{val}'이 있습니다",
                        ))

            # 참조 무결성
            menu_val = rd.get("메뉴명")
            if menu_val and not self._is_blank(menu_val):
                m = str(menu_val).strip()
                if m not in menu_names:
                    issues.append(ValidationIssue(
                        sheet, row_idx, header.get("메뉴명", "A"),
                        f"레시피의 메뉴명 '{m}'이 메뉴 시트에 없습니다",
                    ))

            ing_val = rd.get("식재료명")
            if ing_val and not self._is_blank(ing_val):
                i = str(ing_val).strip()
                if i not in ingredient_names:
                    issues.append(ValidationIssue(
                        sheet, row_idx, header.get("식재료명", "B"),
                        f"레시피의 식재료명 '{i}'이 식재료 시트에 없습니다",
                    ))


# ──────────────────────────────────────────────────────────────
# 적재
# ──────────────────────────────────────────────────────────────

class StoreDataImporter:
    """
    엑셀 파일을 읽어 검증 후 B-06 service로 적재.

    B-06 머지 전: menu_service=None, inventory_service=None → stub 동작.
    B-06 머지 후: app.menu.service, app.inventory.service 인스턴스를 주입.
    """

    def __init__(self, menu_service=None, inventory_service=None) -> None:
        self._menu_svc = menu_service
        self._inventory_svc = inventory_service
        self._validator = StoreDataValidator()

    # ── B-06 stub / 실 service 위임 ───────────────────────────

    def _menu_upsert(self, name: str, price: float, category: str, is_active: bool) -> dict:
        if self._menu_svc:
            return self._menu_svc.upsert(
                name=name, price=price, category=category, is_active=is_active
            )
        return {"id": f"stub_menu_{name}", "name": name}

    def _ingredient_upsert(
        self,
        name: str,
        unit: str,
        unit_price: float,
        current_stock: float,
        safety_stock: float,
        lead_time_days: int,
    ) -> dict:
        if self._inventory_svc:
            return self._inventory_svc.upsert_ingredient(
                name=name,
                unit=unit,
                unit_price=unit_price,
                current_stock=current_stock,
                safety_stock=safety_stock,
                lead_time_days=lead_time_days,
            )
        return {"id": f"stub_ing_{name}", "name": name}

    def _recipe_upsert(
        self, menu_id: str, ingredient_id: str, qty_per_unit: float, unit: str
    ) -> None:
        if self._menu_svc:
            self._menu_svc.upsert_recipe(
                menu_id=menu_id,
                ingredient_id=ingredient_id,
                qty_per_unit=qty_per_unit,
                unit=unit,
            )

    # ── 공개 API ──────────────────────────────────────────────

    def load_file(self, path: str | Path) -> dict:
        """
        엑셀 파일 검증 → 적재.
        오류 시 StoreDataValidationError 발생.
        성공 시 {"menus": N, "ingredients": N, "recipes": N} 반환.
        같은 파일을 두 번 올려도 upsert라 중복 생성 없음.
        """
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        issues = self._validator.validate(wb)
        if issues:
            raise StoreDataValidationError(issues)

        result = {"menus": 0, "ingredients": 0, "recipes": 0}
        menu_id_map: dict[str, str] = {}
        ingredient_id_map: dict[str, str] = {}

        # 메뉴
        ws = wb["메뉴"]
        hdr = self._validator.parse_header(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if self._validator._is_empty_row(row):
                continue
            rd = self._validator._row_dict(row, hdr)
            name = str(rd.get("메뉴명", "")).strip()
            if not name:
                continue
            m = self._menu_upsert(
                name=name,
                price=float(rd.get("가격(원)") or 0),
                category=str(rd.get("카테고리") or "").strip(),
                is_active=str(rd.get("판매중(Y/N)") or "Y").strip().upper() != "N",
            )
            menu_id_map[name] = m["id"]
            result["menus"] += 1

        # 식재료
        ws = wb["식재료"]
        hdr = self._validator.parse_header(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if self._validator._is_empty_row(row):
                continue
            rd = self._validator._row_dict(row, hdr)
            name = str(rd.get("식재료명", "")).strip()
            if not name:
                continue
            i = self._ingredient_upsert(
                name=name,
                unit=str(rd.get("단위") or "").strip(),
                unit_price=float(rd.get("단가(원)") or 0),
                current_stock=float(rd.get("현재고") or 0),
                safety_stock=float(rd.get("안전재고") or 0),
                lead_time_days=int(float(rd.get("리드타임(일)") or 0)),
            )
            ingredient_id_map[name] = i["id"]
            result["ingredients"] += 1

        # 레시피
        ws = wb["레시피"]
        hdr = self._validator.parse_header(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if self._validator._is_empty_row(row):
                continue
            rd = self._validator._row_dict(row, hdr)
            menu_name = str(rd.get("메뉴명", "")).strip()
            ing_name = str(rd.get("식재료명", "")).strip()
            if not menu_name or not ing_name:
                continue
            self._recipe_upsert(
                menu_id=menu_id_map[menu_name],
                ingredient_id=ingredient_id_map[ing_name],
                qty_per_unit=float(rd.get("수량(qty_per_unit)") or 0),
                unit=str(rd.get("단위") or "").strip(),
            )
            result["recipes"] += 1

        return result


# ──────────────────────────────────────────────────────────────
# CLI
# ──────────────────────────────────────────────────────────────

def main() -> None:
    import argparse, json

    ap = argparse.ArgumentParser(description="매장 데이터 엑셀 적재")
    ap.add_argument("xlsx", help="입력 엑셀 파일 경로")
    args = ap.parse_args()

    importer = StoreDataImporter()
    try:
        result = importer.load_file(args.xlsx)
        print(json.dumps(result, ensure_ascii=False))
    except StoreDataValidationError as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
